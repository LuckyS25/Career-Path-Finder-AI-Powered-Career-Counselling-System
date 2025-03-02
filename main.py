import docx
import io
import tkinter as tk
from tkinter import filedialog

def extract_text_from_docx(file_path):
    """Extracts all non-empty paragraphs from a .docx file."""
    doc = docx.Document(file_path)
    return [para.text.strip() for para in doc.paragraphs if para.text.strip()]

def extract_skills(docx_text):
    """
    Extracts skills section based on detecting 'Skills' heading.
    Stops when the next heading (assumed uppercase or short) appears.
    """
    skills = []
    capture = False

    for line in docx_text:
        if "skills" in line.lower():  # Detect 'Skills' heading
            capture = True
            continue

        if capture:
            if line.isupper() or len(line) < 5:  # Assume uppercase or short text = new heading
                break
            skills.append(line)

    return "\n".join(skills).strip()

def extract_skills_from_docx(file_path):
    """
    Extracts skills using document structure (styles like 'Heading').
    Stops extraction when another heading appears.
    """
    doc = docx.Document(file_path)
    skills = []
    skills_found = False

    for para in doc.paragraphs:
        text = para.text.strip()

        # Detect 'Skills' heading based on Word styles
        if para.style.name.startswith("Heading"):
            if skills_found:
                break  # Stop when the next heading appears
            if "skill" in text.lower():
                skills_found = True
                continue

        if skills_found and text:
            skills.append(text)

    return "\n".join(skills).strip()

def main():
    """Main function to select a file and extract skills."""
    root = tk.Tk()
    root.withdraw()  # Hide the main tkinter window

    file_path = filedialog.askopenfilename(title="Select a .docx file", filetypes=[("Word Documents", "*.docx")])

    if not file_path:
        print("No file selected.")
        return

    # Extract text and skills
    docx_text = extract_text_from_docx(file_path)
    skills_text_1 = extract_skills(docx_text)  # Method 1: Text-based
    skills_text_2 = extract_skills_from_docx(file_path)  # Method 2: Style-based

    # Choose the best extraction (fallback to method 2 if method 1 fails)
    extracted_skills = skills_text_1 if skills_text_1 else skills_text_2

    if extracted_skills:
        print("\nExtracted Skills:\n", extracted_skills)
    else:
        print("\nNo skills section found!")

if __name__ == "__main__":
    main()
